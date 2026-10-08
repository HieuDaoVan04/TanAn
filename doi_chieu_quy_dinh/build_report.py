from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib, json, re, math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r'D:\Đồ Án\doi_chieu_quy_dinh')
OUT=Path(r'D:\Đồ Án\Bao_cao_tot_nghiep_Dao_Van_Hieu')
OUT.mkdir(exist_ok=True)
ASSET=ROOT/'report_assets'; ASSET.mkdir(exist_ok=True)
SOURCE=ROOT/'2.Mẫu trang bìa.docx'
REPORT=ROOT/'2025_MauBaoCao.docx'
ASSIGN=Path(r'C:\Users\ADMIN\Downloads\PhieuGiaoDeTai_DaoVanHieu.docx')
NAME='Bao_cao_DATN_Dao_Van_Hieu_2200454'
TITLE='NGHIÊN CỨU VÀ XÂY DỰNG NỀN TẢNG SỐ HỖ TRỢ QUẢN LÝ BIẾN ĐỘNG DÂN CƯ, AN SINH XÃ HỘI VÀ ĐIỀU HÀNH TẠI XÃ TÂN AN TRÊN DỮ LIỆU MÔ PHỎNG'
TITLE_CASE='Nghiên cứu và xây dựng nền tảng số hỗ trợ quản lý biến động dân cư, an sinh xã hội và điều hành tại xã Tân An trên dữ liệu mô phỏng'
contract={'sources':{str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in [SOURCE,REPORT,ASSIGN]},'page':'A4, left 3 cm, other margins 2 cm, header/footer 1.27 cm','body':'Times New Roman Unicode 13, justified, Multiple 1.3','headings':'Black bold; chapter 16 centered new page; level 2 14 and level 3 13; TOC to level 3','cover':'Retained cover table frame and school logo; replace author/title/type/year; compact blank spacing for long real title; secondary cover retains source roles','preserve':'Original source files; original school logo pixels and cover table border pattern','slots':'All sample identities, example titles, dates and report body replaced. Assignment text copied without filling missing dates/signatures. Comments and journal blank for real completion. Original template page numbering replaced by explicit rule starting at acknowledgement. Diagrams authored from inspected code or labeled design.'}
(ROOT/'artifact.md').write_text('# Hợp đồng dựng báo cáo\n\n'+json.dumps(contract,ensure_ascii=False,indent=2),encoding='utf-8')
with ZipFile(SOURCE) as z:
    (ASSET/'logo.png').write_bytes(z.read('word/media/image1.png'))

# Grayscale engineering diagrams; no invented product screenshots.
FONT=Path(r'C:\Windows\Fonts\times.ttf'); BOLD=Path(r'C:\Windows\Fonts\timesbd.ttf')
def diagram(file, nodes, edges, size=(1500,1000)):
    im=Image.new('RGB',size,'white'); dr=ImageDraw.Draw(im)
    font=ImageFont.truetype(str(FONT),30); bold=ImageFont.truetype(str(BOLD),32); small=ImageFont.truetype(str(FONT),25)
    def point(box,side):
        x,y,w,h=box
        return {'l':(x,y+h/2),'r':(x+w,y+h/2),'t':(x+w/2,y),'b':(x+w/2,y+h)}[side]
    for a,b,sa,sb,label in edges:
        p=point(nodes[a][0],sa); q=point(nodes[b][0],sb)
        dr.line([p,q],fill='black',width=3)
        ang=math.atan2(q[1]-p[1],q[0]-p[0]); L=15
        dr.polygon([q,(q[0]-L*math.cos(ang-.45),q[1]-L*math.sin(ang-.45)),(q[0]-L*math.cos(ang+.45),q[1]-L*math.sin(ang+.45))],fill='black')
        if label:
            pos=((p[0]+q[0])/2,(p[1]+q[1])/2-22); bb=dr.textbbox(pos,label,font=small,anchor='mm'); dr.rectangle((bb[0]-5,bb[1]-2,bb[2]+5,bb[3]+2),fill='white'); dr.text(pos,label,font=small,fill='black',anchor='mm')
    for key,(box,title,lines) in nodes.items():
        x,y,w,h=box; dr.rectangle((x,y,x+w,y+h),fill='white',outline='black',width=3)
        dr.rectangle((x+2,y+2,x+w-2,y+53),fill='#eeeeee')
        dr.text((x+w/2,y+28),title,font=bold,fill='black',anchor='mm')
        for j,line in enumerate(lines): dr.text((x+w/2,y+81+j*40),line,font=font,fill='black',anchor='mm')
    im.save(ASSET/file)

diagram('architecture.png',{
'browser':((450,20,600,130),'Trình duyệt',['Cán bộ và người sử dụng']),
'ui':((450,230,600,160),'Blazor Interactive Server',['Giao diện và phiên đăng nhập','Dịch vụ / HTTP quản trị']),
'api':((20,450,410,150),'ASP.NET Core API',['Controller nghiệp vụ','API quản trị']),
'app':((510,470,480,160),'Application',['Dân cư • An sinh • Hồ sơ','Dashboard • Audit • AI']),
'redis':((1110,260,360,150),'Redis',['Phiên làm việc']),
'db':((510,760,480,150),'Infrastructure / EF Core',['Cơ sở dữ liệu quan hệ']),
'ai':((1100,570,380,160),'Dịch vụ AI bên ngoài',['Điểm tích hợp trợ lý','Chưa kiểm chứng thực nghiệm'])
},[('browser','ui','b','t',''),('ui','app','b','t',''),('ui','api','l','t','HTTP'),('api','app','r','l',''),('ui','redis','r','l',''),('app','db','b','t',''),('app','ai','r','l','')],(1500,950))
diagram('erd_population.png',{
'a':((35,30,420,180),'ApThon',['Id • Ma • Ten','GroupId • XaId']),
'h':((540,30,420,180),'HoGiaDinh',['Id • MaSoHo','ApThonId • DiaChi']),
'n':((1045,30,420,180),'NhanKhau',['Id • HoTen • CCCD','MaHoGiaDinh']),
't':((540,360,420,200),'ThanhVienHo',['HoGiaDinhId • NhanKhauId','TuNgay • DenNgay','LaChuHo']),
'b':((1045,360,420,200),'BienDongDanCu',['NhanKhauId','LoaiBienDong','NgayPhatSinh']),
'u':((35,360,420,200),'PhuTrachThon',['UserId','ApThonId','Phân công địa bàn'])
},[('a','h','r','l','1–n'),('h','n','r','l','1–n'),('h','t','b','t','1–n'),('n','t','b','r','1–n'),('n','b','b','t','1–n'),('a','u','b','t','1–n')],(1500,610))
diagram('erd_welfare.png',{
'h':((30,20,420,150),'HoGiaDinh',['Id']),
'n':((540,20,420,150),'NhanKhau',['Id']),
'y':((1045,20,420,150),'YeuCauNguoiDan',['Id • MaYeuCau']),
'p':((30,300,420,180),'PhanLoaiHo',['HoGiaDinhId','TuNgay • DenNgay']),
'd':((540,300,420,180),'DoiTuongAnSinh',['NhanKhauId','Mức trợ cấp']),
'l':((1045,300,420,180),'LichSuXuLyHoSo',['YeuCauId','Diễn biến xử lý']),
't':((540,610,420,150),'LichSuTroCap',['DoiTuongAnSinhId • SoTien']),
'f':((1045,610,420,150),'TepDinhKem',['YeuCauId • TenTep'])
},[('h','p','b','t','1–n'),('n','d','b','t','1–n'),('d','t','b','t','1–n'),('y','l','b','t','1–n'),('y','f','r','r','1–n')],(1530,800))
# Route attachment connection at right margin to avoid passing through another table.
im=Image.open(ASSET/'erd_welfare.png'); dr=ImageDraw.Draw(im); dr.rectangle((1467,25,1525,770),fill='white'); dr.line([(1465,95),(1500,95),(1500,685),(1465,685)],fill='black',width=3); dr.polygon([(1465,685),(1480,677),(1480,693)],fill='black'); dr.text((1500,240),'1–n',font=ImageFont.truetype(str(FONT),25),fill='black',anchor='mm'); im.save(ASSET/'erd_welfare.png')
diagram('head_change.png',{
'a':((210,20,1080,140),'1  Kiểm tra quyền và thành viên',['Người được chọn phải thuộc hộ đang xử lý']),
'b':((210,235,1080,140),'2  Bắt đầu giao dịch',['Kiểm tra giai đoạn hiện hành và ngày hiệu lực']),
'c':((210,450,1080,140),'3  Đóng giai đoạn cũ và mở giai đoạn mới',['Cập nhật quan hệ thành viên và thông tin chủ hộ']),
'd':((210,665,1080,140),'4  Ghi nhật ký và kết thúc giao dịch',['Thành công: xác nhận; lỗi: hủy thay đổi'])
},[('a','b','b','t',''),('b','c','b','t',''),('c','d','b','t','')],(1500,840))

doc=Document(SOURCE)
cover_table=doc.tables[0]
# Keep source frame and source package; rebuild only mapped cover text slots.
for child in list(doc.element.body):
    if child is not cover_table._element and child.tag!=qn('w:sectPr'): doc.element.body.remove(child)

def setfont(run,size=13,bold=None,italic=None):
    run.font.name='Times New Roman'; run.font.size=Pt(size); run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    pr=run._element.get_or_add_rPr(); fonts=pr.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
    for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),'Times New Roman')
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+key),None)

for s in doc.sections:
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(3);s.right_margin=Cm(2);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.header_distance=Cm(1.27);s.footer_distance=Cm(1.27)
    s.different_first_page_header_footer=False
for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption','Title','TOC 1','TOC 2','TOC 3']:
    if name not in doc.styles:
        from docx.enum.style import WD_STYLE_TYPE
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(13);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.3;st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(4)
    st.paragraph_format.widow_control=True
    st.paragraph_format.first_line_indent=Cm(0)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1)
for level,size in [(1,16),(2,14),(3,13)]:
    st=doc.styles['Heading '+str(level)];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10 if level>1 else 0);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.page_break_before=(level==1)
    ppr=st.element.get_or_add_pPr();ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),str(level-1));ppr.append(ol)
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Caption'].paragraph_format.first_line_indent=Cm(0)
doc.styles['Caption'].font.italic=False

def ptext(text='',style=None,align=None,bold=False,size=13,italic=False,container=doc):
    p=container.add_paragraph(style=style or 'Normal');setfont(p.add_run(text),size,bold,italic)
    if align is not None:p.alignment=align;p.paragraph_format.first_line_indent=Cm(0)
    return p
def heading(text,level=1):
    p=doc.add_paragraph(text.replace('\\n',' \n'),style='Heading '+str(level))
    for r in p.runs:setfont(r,{1:16,2:14,3:13}[level],True)
    return p
def field(p,code,placeholder=''):
    r=p.add_run(); b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
    t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+code+' ';r._r.append(t)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    if placeholder:p.add_run(placeholder)
    rr=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');rr._r.append(e)
def plain_title(text):
    p=ptext(text,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,size=16);p.paragraph_format.space_after=Pt(14);p.paragraph_format.keep_with_next=True;return p

cell=cover_table.cell(0,0)
for child in list(cell._tc):
    if child.tag!=qn('w:tcPr'):cell._tc.remove(child)
cover_table.autofit=False;cover_table.alignment=WD_TABLE_ALIGNMENT.CENTER
cover_table.columns[0].width=Cm(16);cell.width=Cm(16)
cover_table.rows[0].height=Cm(24.7);cover_table.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
def cover_p(text,size=14,before=0,after=0,container=cell):
    p=ptext(text,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,size=size,container=container);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after);return p
cover_p('TRƯỜNG ĐẠI HỌC CÔNG NGHIỆP VIỆT – HUNG',14,18)
cover_p('KHOA CÔNG NGHỆ THÔNG TIN',14,2)
lp=cover_p('',before=16);lp.add_run().add_picture(str(ASSET/'logo.png'),width=Cm(3.1))
cover_p('ĐÀO VĂN HIẾU',16,16)
cover_p(TITLE,18,38,18)
cover_p('ĐỒ ÁN TỐT NGHIỆP',18,22)
cover_p('CHUYÊN NGÀNH CÔNG NGHỆ THÔNG TIN',14,6)
cover_p('HÀ NỘI, NĂM 2026',14,85)
# Required paragraph after table also carries page break.
doc.add_page_break()
cover_p('TRƯỜNG ĐẠI HỌC CÔNG NGHIỆP VIỆT – HUNG',14,container=doc)
cover_p('KHOA CÔNG NGHỆ THÔNG TIN',14,4,container=doc)
cover_p(TITLE,18,90,25,container=doc)
cover_p('ĐỒ ÁN TỐT NGHIỆP',18,25,container=doc)
cover_p('CHUYÊN NGÀNH CÔNG NGHỆ THÔNG TIN',14,6,25,container=doc)
for label,value in [('Giảng viên hướng dẫn','ThS. Vũ Hùng Cường'),('Sinh viên thực hiện','Đào Văn Hiếu'),('Mã sinh viên','2200454'),('Lớp','K4628-CNTT'),('Khóa và hệ đào tạo','46 – Chính quy')]:
    p=ptext(label+': '+value,bold=True,size=13,align=WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.left_indent=Cm(1.7);p.paragraph_format.space_after=Pt(5)
cover_p('HÀ NỘI, NĂM 2026',14,75,container=doc)
doc.add_page_break()

# Copy the actual assignment text, including its original blank dates.
assignment=Document(ASSIGN)
count=0;in_products=False
for sp in assignment.paragraphs:
    text=sp.text.strip()
    if not text:continue
    if text=='3. Về sản phẩm':in_products=True;count=0
    elif text.startswith('Thời gian thực hiện:'):in_products=False
    elif in_products:count+=1;text=f'{count}. {text}'
    if 'PHIẾU GIAO ĐỀ TÀI ĐỒ ÁN TỐT NGHIỆP' in text:
        plain_title(text);continue
    if text.startswith('Họ tên sinh viên:'):text='Họ tên sinh viên: Đào Văn Hiếu     Mã SV: 2200454'
    if text.startswith('Khóa/loại hình'):text='Khóa/loại hình đào tạo: 46/Chính quy     Lớp: K4628-CNTT'
    if text.startswith('Ngành:'):text='Ngành: Công nghệ thông tin     Chuyên ngành: Công nghệ thông tin'
    if text.startswith('Điện thoại:'):text='Điện thoại: 0352265243     Email: hieudaovan310@gmail.com'
    if 'TRƯỜNG ĐHCN' in text:
        ptext('TRƯỜNG ĐHCN VIỆT – HUNG',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
        ptext('KHOA CÔNG NGHỆ THÔNG TIN',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
        ptext('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
        ptext('Độc lập – Tự do – Hạnh phúc',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True);continue
    if text.startswith('Khoa: Công nghệ'):continue
    if text.startswith('GIẢNG VIÊN HƯỚNG DẪN'):
        p=ptext('GIẢNG VIÊN HƯỚNG DẪN      TRƯỞNG BỘ MÔN      TRƯỞNG KHOA',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,size=12);p.paragraph_format.keep_with_next=True;continue
    p=ptext(text,bold=bool(re.match(r'^(Kết quả dự kiến|[123](\.[12])?\. Về|1\.[12]\. Kiến thức|Tên đề tài|Mục tiêu đề tài)',text)))
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(3)

doc.add_page_break();plain_title('NHẬN XÉT ĐỒ ÁN TỐT NGHIỆP')
ptext('Chuyên ngành: Công nghệ thông tin',align=WD_ALIGN_PARAGRAPH.CENTER)
ptext('Nhận xét của giảng viên hướng dẫn',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True)
for text in ['Họ tên sinh viên: Đào Văn Hiếu     Mã sinh viên: 2200454','Lớp: K4628-CNTT','Tên đề tài: '+TITLE_CASE,'Người nhận xét: ................................................................................','Đơn vị công tác: ................................................................................']:
    p=ptext(text);p.paragraph_format.first_line_indent=Cm(0)
for title in ['1. Về nội dung và thực hiện nhiệm vụ nghiên cứu','2. Về phương pháp nghiên cứu, độ tin cậy của số liệu','3. Về kết quả của đề tài']:
    ptext(title,bold=True)
    for _ in range(3):ptext('..............................................................................................................')
doc.add_page_break()
for title in ['4. Những thiếu sót và vấn đề cần làm rõ','5. Ý kiến kết luận về mức độ đáp ứng yêu cầu','6. Câu hỏi của người nhận xét']:
    ptext(title,bold=True)
    for _ in range(4):ptext('..............................................................................................................')
ptext('Đánh giá điểm: ................................................................................')
ptext('Kết luận:  □ Đồng ý     □ Không đồng ý cho phép sinh viên bảo vệ.')
ptext('Hà Nội, ngày ...... tháng ...... năm 2026',align=WD_ALIGN_PARAGRAPH.RIGHT,italic=True)
ptext('NGƯỜI NHẬN XÉT',align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)
ptext('(Ký và ghi rõ họ tên)',align=WD_ALIGN_PARAGRAPH.RIGHT,italic=True)

section=doc.add_section(WD_SECTION_START.NEW_PAGE)
section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
for p in section.header.paragraphs:p.clear()
hp=section.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;hp.paragraph_format.first_line_indent=Cm(0);field(hp,'PAGE')
pn=OxmlElement('w:pgNumType');pn.set(qn('w:start'),'1');section._sectPr.append(pn)
plain_title('LỜI CẢM ƠN')
for text in [
'Em xin trân trọng cảm ơn Trường Đại học Công nghiệp Việt–Hung và các thầy cô Khoa Công nghệ thông tin đã giảng dạy, giúp em có nền tảng kiến thức để nghiên cứu và thực hiện đồ án tốt nghiệp.',
'Em xin cảm ơn ThS. Vũ Hùng Cường, giảng viên hướng dẫn đề tài “Nghiên cứu và xây dựng nền tảng số hỗ trợ quản lý biến động dân cư, an sinh xã hội và điều hành tại xã Tân An trên dữ liệu mô phỏng”. Những nội dung của đồ án là dịp để em vận dụng kiến thức về lập trình Web, cơ sở dữ liệu, phân tích thiết kế hệ thống và phương pháp kiểm thử vào một bài toán cụ thể.',
'Trong quá trình hoàn thiện đồ án, em mong nhận được ý kiến nhận xét của giảng viên và hội đồng để tiếp tục sửa chữa những hạn chế về nghiệp vụ, kỹ thuật và cách trình bày. Em xin trân trọng cảm ơn!']:ptext(text)
ptext('Sinh viên thực hiện',align=WD_ALIGN_PARAGRAPH.RIGHT)
ptext('Đào Văn Hiếu',align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[16/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=txt
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,val in [('top',65),('bottom',65),('left',85),('right',85)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.3;p.paragraph_format.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                if ri==0:p.paragraph_format.keep_with_next=True
                for r in p.runs:setfont(r,13,ri==0)
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
    t._tbl.tblPr.append(borders)
    p=ptext('');p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(0);p.runs[0].font.size=Pt(3)
    return t

doc.add_page_break();plain_title('NHẬT KÝ THỰC HIỆN ĐỀ TÀI')
t=table(['Thời gian','Công việc','Ghi chú'],[['','',''] for _ in range(13)],[3,9,4])
for row in t.rows[1:]:row.height=Cm(1.35);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST

doc.add_page_break();plain_title('MỤC LỤC')
field(ptext(''),r'TOC \o "1-3" \h \z \u')
plain_title('DANH MỤC CHỮ VIẾT TẮT').paragraph_format.page_break_before=True
table(['Chữ viết tắt','Tên đầy đủ','Ý nghĩa'],[
['AI','Artificial Intelligence','Trí tuệ nhân tạo'],['API','Application Programming Interface','Giao diện lập trình ứng dụng'],['CCCD','Căn cước công dân','Trường định danh trong mô hình'],['CSDL','Cơ sở dữ liệu','Tập dữ liệu được tổ chức'],['DTO','Data Transfer Object','Đối tượng trao đổi dữ liệu'],['EF Core','Entity Framework Core','Công cụ ánh xạ dữ liệu'],['ERD','Entity Relationship Diagram','Sơ đồ quan hệ thực thể'],['HTTP','Hypertext Transfer Protocol','Giao thức truyền siêu văn bản'],['Id','Identifier','Mã định danh'],['NLP','Natural Language Processing','Xử lý ngôn ngữ tự nhiên'],['RBAC','Role-Based Access Control','Kiểm soát truy cập theo vai trò'],['SDK','Software Development Kit','Bộ công cụ phát triển'],['SQL','Structured Query Language','Ngôn ngữ truy vấn có cấu trúc'],['UI','User Interface','Giao diện người dùng'],['UTC','Coordinated Universal Time','Giờ phối hợp quốc tế']],[2.5,6.5,7])

body=(ROOT/'report_body.md').read_text(encoding='utf-8').splitlines()
tables=[];figures=[]
for line in body:
    if line.startswith('@table '):tables.append(line[7:])
    if line.startswith('@figure '):figures.append(line.split('|',1)[1].strip())
def bookmark(p,name,num):
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(num));b.set(qn('w:name'),name);p._p.insert(0,b)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(num));p._p.append(e)
def catalog(title,items,prefix):
    doc.add_page_break();plain_title(title)
    for i,text in enumerate(items):
        p=ptext(text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(5)
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        p.add_run('\t');field(p,'PAGEREF '+prefix+str(i)+' \\h')
catalog('DANH MỤC BẢNG BIỂU',tables,'tbl')
catalog('DANH MỤC HÌNH VẼ',figures,'fig')

i=0;ti=0;fi=0;in_refs=False
while i<len(body):
    line=body[i].strip();i+=1
    if not line:continue
    if line.startswith('@table '):
        p=ptext(line[7:],style='Caption');p.paragraph_format.keep_with_next=True;bookmark(p,'tbl'+str(ti),100+ti);ti+=1
        head=[x.strip() for x in body[i].split('|')];i+=1;rows=[]
        while i<len(body) and body[i].strip()!='@end':rows.append([x.strip() for x in body[i].split('|')]);i+=1
        i+=1
        if head[0] in ['Mã','Phương thức']:widths=[1.5,6.5,8] if head[0]=='Mã' else [2.2,7.4,6.4]
        elif head[0]=='Bảng và trường':widths=[7.3,3.2,5.5] if head[1]=='Kiểu trong mô hình' else [7.3,3.6,5.1]
        elif line.startswith('@table Bảng 3.2.'):widths=[6.1,3.8,6.1]
        elif len(head)==3:widths=[4.2,5.4,6.4]
        else:widths=None
        table(head,rows,widths);continue
    if line.startswith('@figure '):
        file,caption=[x.strip() for x in line[8:].split('|',1)]
        p=ptext('',align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ASSET/file),width=Cm(15.8))
        p=ptext(caption,style='Caption');bookmark(p,'fig'+str(fi),200+fi);fi+=1;continue
    if line.startswith('@equation '):
        p=ptext(line[10:],align=WD_ALIGN_PARAGRAPH.RIGHT);p.paragraph_format.keep_with_next=True;continue
    if line.startswith('# '):
        heading(line[2:],1);in_refs=line[2:]=='TÀI LIỆU THAM KHẢO';continue
    if line.startswith('## '):heading(line[3:],2);continue
    if line.startswith('### '):heading(line[4:],3);continue
    if line.startswith('~'):
        p=ptext('Tóm tắt chương',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True);p.paragraph_format.keep_with_next=True
        p=ptext(line[1:],italic=True);p.paragraph_format.space_after=Pt(10);continue
    p=ptext(line)
    if in_refs:
        p.paragraph_format.left_indent=Cm(1);p.paragraph_format.first_line_indent=Cm(-1)
        # Titles are italicized as required by the reference convention.
        match=re.match(r'(\[\d+\] [^,]+, )([^,]+)(.*)',line)
        if match:
            p.clear();setfont(p.add_run(match[1]));setfont(p.add_run(match[2]),italic=True);setfont(p.add_run(match[3]))

# Normalize headers, footers and inherited cover section flags.
for idx,s in enumerate(doc.sections):
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(3);s.right_margin=Cm(2);s.header_distance=Cm(1.27);s.footer_distance=Cm(1.27)
    for p in s.footer.paragraphs:p.clear()
    if idx==0:
        for p in s.header.paragraphs:p.clear()
    for p in s.header.paragraphs:
        for r in p.runs:setfont(r,13)
settings=doc.settings.element; uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
doc.core_properties.title=TITLE_CASE;doc.core_properties.author='Đào Văn Hiếu';doc.core_properties.subject='Đồ án tốt nghiệp Công nghệ thông tin';doc.core_properties.comments=''
dest=OUT/(NAME+'.docx');doc.save(dest)
print(json.dumps({'output':str(dest),'tables':ti,'figures':fi,'body_words':len(' '.join(body).split()),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
